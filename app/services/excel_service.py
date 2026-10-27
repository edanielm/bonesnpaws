"""
app/services/excel_service.py

Reportes descargables en Excel (.xlsx):

    compras_cliente(id_usuario)  -> "Mis compras": las facturas de un
                                    cliente y lo que compro en cada una.
    reporte_ventas()             -> reporte completo del panel admin:
                                    resumen, facturas, detalle, productos y
                                    clientes, una hoja por tema.

Se usa openpyxl, Python puro como ReportLab y python-docx. Los montos se
escriben como numeros (con formato de pesos), no como texto: asi el archivo
se puede sumar, filtrar y graficar en Excel sin convertir nada.
"""
from datetime import date
from io import BytesIO

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload, selectinload

from app.formato import folio
from app.models import (
    Categoria, Cliente, Detalle, Factura, Producto, Proveedor, Transportadora,
)
from app.services.admin_service import AdminService

FORMATO_PESOS = '"$"#,##0'
FORMATO_FECHA = "DD/MM/YYYY"

ENCABEZADO_FONDO = PatternFill("solid", fgColor="4D86B8")
ENCABEZADO_LETRA = Font(bold=True, color="FFFFFF")
TITULO_LETRA = Font(bold=True, size=14, color="4D86B8")


def _f(valor):
    """Numeric/Decimal de la base -> float, para que Excel lo trate como numero."""
    return float(valor) if valor is not None else 0.0


def _hoja(libro, titulo, columnas, filas):
    """
    Agrega una hoja con encabezado, filtros y la primera fila congelada.

    `columnas` es una lista de (titulo, ancho, formato); formato es None,
    FORMATO_PESOS o FORMATO_FECHA y se aplica a toda la columna.
    """
    hoja = libro.create_sheet(titulo)
    hoja.append([c[0] for c in columnas])
    for celda in hoja[1]:
        celda.fill = ENCABEZADO_FONDO
        celda.font = ENCABEZADO_LETRA
        celda.alignment = Alignment(vertical="center")

    for fila in filas:
        hoja.append(fila)

    for i, (_, ancho, formato) in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        hoja.column_dimensions[letra].width = ancho
        if formato:
            for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato

    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = hoja.dimensions
    return hoja


def _guardar(libro):
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _facturas_con_detalle(consulta):
    """Facturas con su detalle, productos y envio cargados en pocas consultas."""
    return (
        consulta.options(
            selectinload(Factura.detalles).joinedload(Detalle.producto),
            joinedload(Factura.envio),
        )
        .order_by(Factura.fecha.desc(), Factura.numero_factura.desc())
        .all()
    )


def _nombre_producto(detalle):
    return detalle.producto.nombre if detalle.producto else "Producto retirado del catálogo"


class ExcelService:

    # ------------------------------------------------------------------
    # Cliente: "Mis compras"
    # ------------------------------------------------------------------

    @staticmethod
    def compras_cliente(id_usuario):
        """Devuelve (bytes_del_xlsx, nombre_de_archivo) con las compras del cliente."""
        prefijo = current_app.config["FACTURA_PREFIJO"]
        facturas = _facturas_con_detalle(Factura.query.filter_by(id_usuario=id_usuario))

        libro = Workbook()
        libro.remove(libro.active)

        _hoja(libro, "Mis compras", [
            ("Factura", 13, None),
            ("Fecha", 12, FORMATO_FECHA),
            ("Método de pago", 16, None),
            ("Productos", 11, None),
            ("Subtotal", 14, FORMATO_PESOS),
            ("IVA", 12, FORMATO_PESOS),
            ("Envío", 12, FORMATO_PESOS),
            ("Total", 14, FORMATO_PESOS),
            ("Estado del envío", 17, None),
            ("Guía", 18, None),
        ], [
            [
                folio(f.numero_factura, prefijo), f.fecha, f.metodo_pago,
                sum(d.cantidad for d in f.detalles),
                _f(f.subtotal), _f(f.iva), _f(f.costo_envio), _f(f.total),
                f.envio.estado if f.envio else "Sin envío",
                f.envio.numero_guia if f.envio else "",
            ]
            for f in facturas
        ])

        _hoja(libro, "Productos comprados", [
            ("Factura", 13, None),
            ("Fecha", 12, FORMATO_FECHA),
            ("Código", 12, None),
            ("Producto", 42, None),
            ("Cantidad", 10, None),
            ("Precio unitario", 15, FORMATO_PESOS),
            ("Subtotal", 14, FORMATO_PESOS),
        ], [
            [
                folio(f.numero_factura, prefijo), f.fecha,
                d.producto.codigo if d.producto else "", _nombre_producto(d),
                d.cantidad, _f(d.precio_unitario), _f(d.subtotal_producto),
            ]
            for f in facturas for d in f.detalles
        ])

        return _guardar(libro), f"mis-compras-{date.today():%Y%m%d}.xlsx"

    # ------------------------------------------------------------------
    # Administrador: reporte completo de ventas
    # ------------------------------------------------------------------

    @staticmethod
    def reporte_ventas():
        """Devuelve (bytes_del_xlsx, nombre_de_archivo) con el reporte del panel."""
        prefijo = current_app.config["FACTURA_PREFIJO"]
        facturas = _facturas_con_detalle(Factura.query)
        clientes = {c.id_usuario: c for c in Cliente.query.all()}
        transportadoras = {t.id_transportadora: t.nombre for t in Transportadora.query.all()}
        categorias = {c.id_categoria: c.nombre for c in Categoria.query.all()}
        proveedores = {p.id_proveedor: p.nombre for p in Proveedor.query.all()}

        libro = Workbook()
        libro.remove(libro.active)

        ExcelService._hoja_resumen(libro)

        def nombre_cliente(id_usuario):
            c = clientes.get(id_usuario)
            return f"{c.nombre} {c.apellido}" if c else id_usuario

        _hoja(libro, "Facturas", [
            ("Factura", 13, None),
            ("Fecha", 12, FORMATO_FECHA),
            ("Documento", 13, None),
            ("Cliente", 26, None),
            ("Correo", 28, None),
            ("Teléfono", 14, None),
            ("Dirección", 30, None),
            ("Método de pago", 16, None),
            ("Unidades", 10, None),
            ("Subtotal", 14, FORMATO_PESOS),
            ("IVA", 12, FORMATO_PESOS),
            ("Envío", 12, FORMATO_PESOS),
            ("Total", 14, FORMATO_PESOS),
            ("Estado del envío", 17, None),
            ("Transportadora", 16, None),
            ("Guía", 18, None),
        ], [
            [
                folio(f.numero_factura, prefijo), f.fecha, f.id_usuario,
                nombre_cliente(f.id_usuario),
                clientes[f.id_usuario].correo if f.id_usuario in clientes else "",
                clientes[f.id_usuario].telefono if f.id_usuario in clientes else "",
                clientes[f.id_usuario].direccion if f.id_usuario in clientes else "",
                f.metodo_pago, sum(d.cantidad for d in f.detalles),
                _f(f.subtotal), _f(f.iva), _f(f.costo_envio), _f(f.total),
                f.envio.estado if f.envio else "Sin envío",
                transportadoras.get(f.envio.id_transportadora, "") if f.envio else "",
                f.envio.numero_guia if f.envio else "",
            ]
            for f in facturas
        ])

        _hoja(libro, "Detalle de ventas", [
            ("Factura", 13, None),
            ("Fecha", 12, FORMATO_FECHA),
            ("Cliente", 26, None),
            ("Código", 12, None),
            ("Producto", 42, None),
            ("Categoría", 14, None),
            ("Cantidad", 10, None),
            ("Precio unitario", 15, FORMATO_PESOS),
            ("Subtotal", 14, FORMATO_PESOS),
        ], [
            [
                folio(f.numero_factura, prefijo), f.fecha, nombre_cliente(f.id_usuario),
                d.producto.codigo if d.producto else "", _nombre_producto(d),
                categorias.get(d.producto.id_categoria, "") if d.producto else "",
                d.cantidad, _f(d.precio_unitario), _f(d.subtotal_producto),
            ]
            for f in facturas for d in f.detalles
        ])

        # Lo vendido por producto se suma aqui mismo y no con una consulta
        # aparte: el detalle ya esta en memoria.
        vendido = {}
        for f in facturas:
            for d in f.detalles:
                unidades, ingresos = vendido.get(d.id_producto, (0, 0.0))
                vendido[d.id_producto] = (unidades + d.cantidad, ingresos + _f(d.subtotal_producto))

        _hoja(libro, "Productos", [
            ("Código", 12, None),
            ("Producto", 42, None),
            ("Categoría", 14, None),
            ("Proveedor", 24, None),
            ("Precio", 14, FORMATO_PESOS),
            ("Stock", 9, None),
            ("Unidades vendidas", 18, None),
            ("Ingresos", 15, FORMATO_PESOS),
            ("Estado", 11, None),
        ], [
            [
                p.codigo, p.nombre, categorias.get(p.id_categoria, ""),
                proveedores.get(p.id_proveedor, ""), _f(p.precio), p.stock,
                *vendido.get(p.id_producto, (0, 0.0)),
                "Activo" if p.activo else "Retirado",
            ]
            for p in Producto.query.order_by(Producto.nombre).all()
        ])

        compras = {}
        for f in facturas:
            num, gastado = compras.get(f.id_usuario, (0, 0.0))
            compras[f.id_usuario] = (num + 1, gastado + _f(f.total))

        _hoja(libro, "Clientes", [
            ("Documento", 13, None),
            ("Nombre", 26, None),
            ("Correo", 28, None),
            ("Teléfono", 14, None),
            ("Dirección", 30, None),
            ("Compras", 10, None),
            ("Total comprado", 16, FORMATO_PESOS),
        ], [
            [
                c.id_usuario, f"{c.nombre} {c.apellido}", c.correo, c.telefono,
                c.direccion, *compras.get(c.id_usuario, (0, 0.0)),
            ]
            for c in sorted(clientes.values(), key=lambda c: (c.nombre, c.apellido))
        ])

        return _guardar(libro), f"reporte-ventas-{date.today():%Y%m%d}.xlsx"

    @staticmethod
    def _hoja_resumen(libro):
        """Primera hoja: los mismos indicadores que muestra el panel."""
        resumen = AdminService.resumen_ventas()
        hoja = libro.create_sheet("Resumen")
        hoja.column_dimensions["A"].width = 30
        hoja.column_dimensions["B"].width = 18
        hoja.column_dimensions["C"].width = 18

        hoja["A1"] = "Reporte de ventas — Bones & Paws"
        hoja["A1"].font = TITULO_LETRA
        hoja["A2"] = "Generado el"
        hoja["B2"] = date.today()
        hoja["B2"].number_format = FORMATO_FECHA

        fila = 4
        for etiqueta, valor, formato in (
            ("Ventas totales", resumen["total_ventas"], FORMATO_PESOS),
            ("Facturas", resumen["num_facturas"], None),
            ("Ticket promedio", resumen["ticket_promedio"], FORMATO_PESOS),
            ("Unidades vendidas", resumen["unidades_vendidas"], None),
            ("Ventas del mes", resumen["ventas_mes"], FORMATO_PESOS),
            ("Ventas del mes anterior", resumen["ventas_mes_anterior"], FORMATO_PESOS),
            ("Clientes registrados", resumen["num_clientes"], None),
            ("Productos", resumen["num_productos"], None),
            ("Productos agotados", resumen["productos_agotados"], None),
            ("Productos con stock bajo (1 a 5)", resumen["productos_stock_bajo"], None),
        ):
            hoja.cell(fila, 1, etiqueta).font = Font(bold=True)
            celda = hoja.cell(fila, 2, valor)
            if formato:
                celda.number_format = formato
            fila += 1

        fila += 1
        for titulo, columnas, datos in (
            ("Ventas por mes (últimos 12)", ("Mes", "Facturas", "Total"),
             [(f"{p['etiqueta']} {p['anio']}", p["facturas"], p["total"])
              for p in AdminService.ventas_por_mes(12)]),
            ("Ventas por método de pago", ("Método", "Facturas", "Total"),
             [(m["metodo"], m["facturas"], m["total"])
              for m in AdminService.ventas_por_metodo_pago()]),
            ("Productos más vendidos", ("Producto", "Unidades", "Ingresos"),
             [(p["nombre"], p["unidades"], p["ingresos"])
              for p in AdminService.top_productos(10)]),
        ):
            hoja.cell(fila, 1, titulo).font = Font(bold=True, color="4D86B8")
            fila += 1
            for col, encabezado in enumerate(columnas, start=1):
                celda = hoja.cell(fila, col, encabezado)
                celda.fill = ENCABEZADO_FONDO
                celda.font = ENCABEZADO_LETRA
            fila += 1
            for nombre, cantidad, total in datos:
                hoja.cell(fila, 1, nombre)
                hoja.cell(fila, 2, cantidad)
                hoja.cell(fila, 3, total).number_format = FORMATO_PESOS
                fila += 1
            fila += 1
