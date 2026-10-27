"""
app/services/factura_word_service.py

Genera la factura de venta como documento de Word (.docx).

Es la misma factura que arma FacturaPdfService -mismos datos, mismo orden y
misma paleta-, pero en un formato que el cliente puede abrir y editar en
Word, LibreOffice o Google Docs. Se usa python-docx, que igual que ReportLab
es Python puro y se instala con pip en cualquier sistema operativo.

Como el PDF, el documento se arma en memoria y se devuelve como bytes.
"""
import os
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.formato import pesos, pesos_letras, folio
from app.services.factura_pdf_service import _fecha_larga

# Misma paleta del PDF (y del sitio), en el formato que pide Word
AZUL = "4d86b8"
AZUL_SUAVE = "e2ecf6"
ROSA = "e4718f"
TEXTO = "1a2c3f"
TEXTO_TENUE = "7a8c9e"
FILA_ALTERNA = "f7fafd"


def _color(hexa):
    return RGBColor.from_string(hexa.upper())


def _texto(parrafo, contenido, tamano=9, negrita=False, color=TEXTO):
    """Agrega un tramo de texto con formato a un parrafo y lo devuelve."""
    tramo = parrafo.add_run(contenido)
    tramo.font.size = Pt(tamano)
    tramo.font.bold = negrita
    tramo.font.color.rgb = _color(color)
    return tramo


def _parrafo(celda_o_doc, alineacion=None, espacio_despues=0):
    parrafo = celda_o_doc.add_paragraph()
    parrafo.paragraph_format.space_after = Pt(espacio_despues)
    parrafo.paragraph_format.space_before = Pt(0)
    if alineacion is not None:
        parrafo.alignment = alineacion
    return parrafo


def _primer_parrafo(celda, alineacion=None):
    """Una celda nueva ya trae un parrafo vacio: se usa en vez de agregar otro."""
    parrafo = celda.paragraphs[0]
    parrafo.paragraph_format.space_after = Pt(0)
    if alineacion is not None:
        parrafo.alignment = alineacion
    return parrafo


def _sombrear(celda, hexa):
    """Color de fondo de una celda (python-docx no lo expone directamente)."""
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), hexa)
    celda._tc.get_or_add_tcPr().append(sombra)


def _sin_bordes(tabla):
    """Quita los bordes de una tabla usada solo para maquetar."""
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "nil")
        bordes.append(borde)
    tabla._tbl.tblPr.append(bordes)


def _linea_inferior(celda, hexa, grosor=4):
    """Borde inferior de una celda: separa las filas del detalle."""
    bordes = OxmlElement("w:tcBorders")
    borde = OxmlElement("w:bottom")
    borde.set(qn("w:val"), "single")
    borde.set(qn("w:sz"), str(grosor))
    borde.set(qn("w:color"), hexa)
    bordes.append(borde)
    celda._tc.get_or_add_tcPr().append(bordes)


def _anchos(tabla, anchos_mm):
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_mm):
            celda.width = Mm(ancho)


class FacturaWordService:

    @staticmethod
    def generar(factura, cliente, empresa, prefijo="BP", nota="", logo=None):
        """
        Devuelve (bytes_del_docx, nombre_de_archivo).

        Recibe lo mismo que FacturaPdfService.generar y con las mismas
        reglas: `factura` trae sus `detalles` cargados y, si el logo no se
        puede leer, el documento sale igual sin el.
        """
        numero = folio(factura.numero_factura, prefijo)
        doc = Document()

        seccion = doc.sections[0]
        seccion.page_width, seccion.page_height = Mm(215.9), Mm(279.4)  # carta
        seccion.left_margin = seccion.right_margin = Mm(18)
        seccion.top_margin = seccion.bottom_margin = Mm(16)
        ancho_util = 215.9 - 36

        doc.core_properties.title = f"Factura {numero} - {empresa['nombre']}"
        doc.core_properties.author = empresa["nombre"]
        doc.core_properties.subject = "Factura de venta"

        estilo = doc.styles["Normal"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(9)

        # ---------- Encabezado: vendedor a la izquierda, factura a la derecha ----------
        encabezado = doc.add_table(rows=1, cols=3)
        _sin_bordes(encabezado)
        _anchos(encabezado, [20, ancho_util - 82, 62])
        celda_logo, celda_emisor, celda_factura = encabezado.rows[0].cells

        if logo and os.path.isfile(logo):
            try:
                _primer_parrafo(celda_logo).add_run().add_picture(logo, width=Mm(16))
            except Exception:
                pass  # sin logo la factura sigue siendo valida

        _texto(_primer_parrafo(celda_emisor), empresa["nombre"], 16, True, AZUL)
        for linea in (
            f"NIT {empresa['nit']}",
            empresa["direccion"],
            empresa["ciudad"],
            f"Tel. {empresa['telefono']} · {empresa['correo']}",
        ):
            _texto(_parrafo(celda_emisor), linea, 8.5, color=TEXTO_TENUE)

        derecha = WD_ALIGN_PARAGRAPH.RIGHT
        _texto(_primer_parrafo(celda_factura, derecha), "FACTURA DE VENTA", 12, True)
        _texto(_parrafo(celda_factura, derecha), numero, 15, True, ROSA)
        _texto(_parrafo(celda_factura, derecha), "Fecha de expedición", 8.5, color=TEXTO_TENUE)
        _texto(_parrafo(celda_factura, derecha), _fecha_larga(factura.fecha), 8.5, color=TEXTO_TENUE)

        # Franja azul bajo el encabezado, como en el PDF
        franja = doc.add_table(rows=1, cols=1)
        _sin_bordes(franja)
        _sombrear(franja.rows[0].cells[0], AZUL)
        franja.rows[0].height = Mm(0.8)
        _primer_parrafo(franja.rows[0].cells[0]).paragraph_format.line_spacing = Pt(1)
        _parrafo(doc, espacio_despues=4)

        # ---------- Datos del cliente ----------
        bloque = doc.add_table(rows=2, cols=2)
        _sin_bordes(bloque)
        _anchos(bloque, [ancho_util * 0.52, ancho_util * 0.48])
        for fila in bloque.rows:
            for celda in fila.cells:
                _sombrear(celda, AZUL_SUAVE)
        titulo = bloque.cell(0, 0).merge(bloque.cell(0, 1))
        _texto(_primer_parrafo(titulo), "FACTURAR A", 8, True, AZUL)

        izquierda, derecha_cliente = bloque.rows[1].cells
        _texto(_primer_parrafo(izquierda), f"{cliente.nombre} {cliente.apellido}", 9, True)
        _texto(_parrafo(izquierda), f"Documento: {cliente.id_usuario}")
        _texto(_parrafo(izquierda), cliente.correo)

        _texto(_primer_parrafo(derecha_cliente), f"Teléfono: {cliente.telefono or '-'}")
        _texto(_parrafo(derecha_cliente), f"Dirección: {cliente.direccion or '-'}")
        pago = _parrafo(derecha_cliente)
        _texto(pago, "Forma de pago: ")
        _texto(pago, factura.metodo_pago, negrita=True)
        _parrafo(doc, espacio_despues=4)

        # ---------- Detalle de productos ----------
        columnas = ["#", "Descripción", "Cant.", "Valor unitario", "Valor total"]
        alineaciones = [WD_ALIGN_PARAGRAPH.CENTER, None, WD_ALIGN_PARAGRAPH.CENTER, derecha, derecha]
        tabla = doc.add_table(rows=1, cols=len(columnas))
        _sin_bordes(tabla)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

        for celda, titulo_col, alineacion in zip(tabla.rows[0].cells, columnas, alineaciones):
            _sombrear(celda, AZUL)
            _texto(_primer_parrafo(celda, alineacion), titulo_col, 8.5, True, "ffffff")

        for i, d in enumerate(factura.detalles, start=1):
            nombre = d.producto.nombre if d.producto else "Producto retirado del catálogo"
            valores = [str(i), nombre, str(d.cantidad), pesos(d.precio_unitario), pesos(d.subtotal_producto)]
            celdas = tabla.add_row().cells
            for celda, valor, alineacion in zip(celdas, valores, alineaciones):
                _texto(_primer_parrafo(celda, alineacion), valor, 8.5)
                _linea_inferior(celda, "d3e0ec")
                if i % 2 == 0:
                    _sombrear(celda, FILA_ALTERNA)

        _anchos(tabla, [10, ancho_util - 86, 16, 30, 30])
        _parrafo(doc, espacio_despues=4)

        # ---------- Monto en letras + totales ----------
        cierre = doc.add_table(rows=1, cols=2)
        _sin_bordes(cierre)
        _anchos(cierre, [ancho_util - 64, 64])
        celda_letras, celda_totales = cierre.rows[0].cells

        letras = _primer_parrafo(celda_letras)
        _texto(letras, "Son: ", 8.5, color=TEXTO_TENUE)
        _texto(letras, pesos_letras(factura.total), 8.5)

        totales = celda_totales.add_table(rows=0, cols=2)
        _sin_bordes(totales)
        for etiqueta, valor in (
            ("Subtotal", factura.subtotal),
            ("IVA (19%)", factura.iva),
            ("Envío", factura.costo_envio),
        ):
            fila = totales.add_row().cells
            _texto(_primer_parrafo(fila[0]), etiqueta, 9, color=TEXTO_TENUE)
            _texto(_primer_parrafo(fila[1], derecha), pesos(valor), 9, color=TEXTO_TENUE)
        fila = totales.add_row().cells
        _texto(_primer_parrafo(fila[0]), "TOTAL A PAGAR", 10.5, True, ROSA)
        _texto(_primer_parrafo(fila[1], derecha), pesos(factura.total), 10.5, True, ROSA)
        _anchos(totales, [32, 30])
        # La celda contenedora ya trae un parrafo vacio antes de la tabla anidada
        celda_totales._tc.remove(celda_totales.paragraphs[0]._p)
        _parrafo(doc, espacio_despues=10)

        # ---------- Pie ----------
        centro = WD_ALIGN_PARAGRAPH.CENTER
        _texto(_parrafo(doc, centro, 3),
               f"Gracias por comprar en {empresa['nombre']} · {empresa['web']}", 7.5, color=TEXTO_TENUE)
        if nota:
            _texto(_parrafo(doc, centro), nota, 7.5, color=TEXTO_TENUE)

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue(), f"factura-{numero}.docx"
